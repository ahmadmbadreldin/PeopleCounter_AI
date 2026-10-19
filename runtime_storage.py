"""SQLite state/event storage, daily hourly Excel reports, and heartbeat files."""

import json
import logging
import os
import sqlite3
import threading
import time
from datetime import date, datetime, timedelta
from pathlib import Path


STATE_KEYS = (
    "total_in",
    "total_out",
    "staff_ignored",
    "unknown_crossings",
    "peak_inside",
)
LOGGER = logging.getLogger("people_counter.storage")


def empty_state(business_date=None):
    return {
        "business_date": business_date or date.today().isoformat(),
        "total_in": 0,
        "total_out": 0,
        "staff_ignored": 0,
        "unknown_crossings": 0,
        "peak_inside": 0,
    }


class SQLiteStore:
    """Small, durable operational store. Excel is generated only as a report."""

    def __init__(
        self,
        database,
        branch_id,
        camera_ip,
        reports_dir,
        event_retention_days=90,
        report_retention_days=365,
        report_refresh_delay_s=2.0,
    ):
        self.database = Path(database)
        self.reports_dir = Path(reports_dir)
        self.branch_id = str(branch_id)
        self.camera_ip = str(camera_ip)
        self.event_retention_days = max(7, int(event_retention_days))
        self.report_retention_days = max(7, int(report_retention_days))
        self.report_refresh_delay_s = max(0.25, float(report_refresh_delay_s))
        self.database.parent.mkdir(parents=True, exist_ok=True)
        self.reports_dir.mkdir(parents=True, exist_ok=True)
        self._lock = threading.RLock()
        self._closed = False
        self.connection = sqlite3.connect(
            str(self.database), timeout=15.0, check_same_thread=False
        )
        self.connection.row_factory = sqlite3.Row
        with self._lock:
            self.connection.execute("PRAGMA journal_mode=WAL")
            self.connection.execute("PRAGMA synchronous=NORMAL")
            self.connection.execute("PRAGMA busy_timeout=15000")
            self._create_schema()
        self._report_requested = threading.Event()
        self._report_stop = threading.Event()
        self._export_lock = threading.Lock()
        self._report_thread = threading.Thread(
            target=self._report_loop,
            name="DailyExcelReport",
            daemon=True,
        )
        self._report_thread.start()

    def _create_schema(self):
        self.connection.executescript(
            """
            CREATE TABLE IF NOT EXISTS counter_state (
                branch_id TEXT PRIMARY KEY,
                business_date TEXT NOT NULL,
                total_in INTEGER NOT NULL DEFAULT 0,
                total_out INTEGER NOT NULL DEFAULT 0,
                staff_ignored INTEGER NOT NULL DEFAULT 0,
                unknown_crossings INTEGER NOT NULL DEFAULT 0,
                peak_inside INTEGER NOT NULL DEFAULT 0,
                updated_at TEXT NOT NULL
            );
            CREATE TABLE IF NOT EXISTS events (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                timestamp TEXT NOT NULL,
                business_date TEXT NOT NULL,
                branch_id TEXT NOT NULL,
                camera_ip TEXT NOT NULL,
                track_id INTEGER NOT NULL,
                event TEXT NOT NULL,
                entered INTEGER NOT NULL,
                exited INTEGER NOT NULL,
                inside INTEGER NOT NULL
            );
            CREATE INDEX IF NOT EXISTS idx_events_branch_date
                ON events(branch_id, business_date, id);
            CREATE TABLE IF NOT EXISTS health_events (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                timestamp TEXT NOT NULL,
                business_date TEXT NOT NULL,
                branch_id TEXT NOT NULL,
                event TEXT NOT NULL,
                details TEXT NOT NULL DEFAULT ''
            );
            CREATE INDEX IF NOT EXISTS idx_health_branch_date
                ON health_events(branch_id, business_date, id);
            """
        )
        self.connection.commit()

    @staticmethod
    def _now():
        return datetime.now().astimezone().isoformat(timespec="seconds")

    def _state_from_row(self, row):
        state = empty_state(row["business_date"])
        for key in STATE_KEYS:
            state[key] = int(row[key])
        return state

    def prepare_business_day(self, business_date=None):
        today = business_date or date.today().isoformat()
        previous_date = None
        with self._lock:
            row = self.connection.execute(
                "SELECT * FROM counter_state WHERE branch_id = ?",
                (self.branch_id,),
            ).fetchone()
            if row is None:
                state = empty_state(today)
                self._write_state_locked(state)
                self.connection.commit()
            elif row["business_date"] == today:
                state = self._state_from_row(row)
            else:
                previous_date = row["business_date"]
                state = empty_state(today)
                self._write_state_locked(state)
                self.connection.commit()
        if previous_date:
            self.export_daily_excel(previous_date)
        self.export_missing_past_reports(today)
        return state

    def _write_state_locked(self, state):
        self.connection.execute(
            """
            INSERT INTO counter_state (
                branch_id, business_date, total_in, total_out,
                staff_ignored, unknown_crossings, peak_inside, updated_at
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?)
            ON CONFLICT(branch_id) DO UPDATE SET
                business_date=excluded.business_date,
                total_in=excluded.total_in,
                total_out=excluded.total_out,
                staff_ignored=excluded.staff_ignored,
                unknown_crossings=excluded.unknown_crossings,
                peak_inside=excluded.peak_inside,
                updated_at=excluded.updated_at
            """,
            (
                self.branch_id,
                state["business_date"],
                int(state["total_in"]),
                int(state["total_out"]),
                int(state["staff_ignored"]),
                int(state["unknown_crossings"]),
                int(state["peak_inside"]),
                self._now(),
            ),
        )
    def save_state(self, state):
        with self._lock:
            self._write_state_locked(state)
            self.connection.commit()

    def add_event(self, track_id, event, state):
        now = self._now()
        inside = max(0, int(state["total_in"]) - int(state["total_out"]))
        with self._lock:
            self.connection.execute("BEGIN IMMEDIATE")
            try:
                self.connection.execute(
                    """
                    INSERT INTO events (
                        timestamp, business_date, branch_id, camera_ip,
                        track_id, event, entered, exited, inside
                    ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
                    """,
                    (
                        now,
                        state["business_date"],
                        self.branch_id,
                        self.camera_ip,
                        int(track_id),
                        str(event),
                        int(state["total_in"]),
                        int(state["total_out"]),
                        inside,
                    ),
                )
                self._write_state_locked(state)
            except Exception:
                self.connection.rollback()
                raise
            else:
                self.connection.commit()
        self.request_report()

    def log_health(self, event, details="", business_date=None):
        with self._lock:
            self.connection.execute(
                """
                INSERT INTO health_events (
                    timestamp, business_date, branch_id, event, details
                ) VALUES (?, ?, ?, ?, ?)
                """,
                (
                    self._now(),
                    business_date or date.today().isoformat(),
                    self.branch_id,
                    str(event),
                    str(details)[:2000],
                ),
            )
            self.connection.commit()
        self.request_report()

    def request_report(self):
        """Refresh today's workbook soon without blocking video processing."""
        if not self._closed:
            self._report_requested.set()

    def _report_loop(self):
        last_periodic_refresh = time.monotonic()
        while not self._report_stop.is_set():
            requested = self._report_requested.wait(timeout=1.0)
            periodic = time.monotonic() - last_periodic_refresh >= 60.0
            if not requested and not periodic:
                continue
            if self._report_stop.is_set():
                break
            # Group events that happen within a short burst into one workbook save.
            if requested:
                time.sleep(self.report_refresh_delay_s)
            if self._report_stop.is_set():
                break
            self._report_requested.clear()
            try:
                self.export_daily_excel(date.today().isoformat())
                last_periodic_refresh = time.monotonic()
            except Exception:
                LOGGER.exception("Unable to refresh the daily Excel report")
                # Excel may temporarily lock the workbook while a user has it
                # open. Keep SQLite authoritative and retry in the background.
                self._report_requested.set()
                self._report_stop.wait(timeout=30.0)

    def export_missing_past_reports(self, today=None):
        today = today or date.today().isoformat()
        with self._lock:
            rows = self.connection.execute(
                """
                SELECT business_date FROM events
                WHERE branch_id = ? AND business_date < ?
                UNION
                SELECT business_date FROM health_events
                WHERE branch_id = ? AND business_date < ?
                """,
                (self.branch_id, today, self.branch_id, today),
            ).fetchall()
        days = sorted({row["business_date"] for row in rows})
        for day in days:
            report = self.reports_dir / f"{day}.xlsx"
            if not report.exists():
                self.export_daily_excel(day)

    def export_monthly_excel(self, report_month=None):
        """Backward-compatible alias; current reports are daily, not monthly."""
        value = str(report_month or date.today().isoformat())
        if len(value) == 7:
            today = date.today().isoformat()
            value = today if today.startswith(value) else f"{value}-01"
        return self.export_daily_excel(value)

    def export_daily_excel(self, business_date=None):
        with self._export_lock:
            return self._export_daily_excel_unlocked(business_date)

    @staticmethod
    def _duration_seconds(rows, start_event, stop_events, report_day, now):
        """Sum event-state intervals for one day, including an open interval."""
        start = None
        total = 0.0
        for row in rows:
            event = row["event"]
            timestamp = datetime.fromisoformat(row["timestamp"])
            if event == start_event:
                if start is None:
                    start = timestamp
            elif event in stop_events and start is not None:
                total += max(0.0, (timestamp - start).total_seconds())
                start = None
        if start is not None:
            if report_day == now.date().isoformat():
                end = now
            else:
                end = start.replace(hour=23, minute=59, second=59)
            total += max(0.0, (end - start).total_seconds())
        return int(round(total))

    @staticmethod
    def _duration_seconds_between(
        rows, start_event, stop_events, interval_start, interval_end, now
    ):
        """Duration of one health state clipped to a single hourly interval."""
        effective_end = min(interval_end, now)
        if interval_start >= effective_end:
            return 0
        active = False
        cursor = interval_start
        total = 0.0
        for row in rows:
            timestamp = datetime.fromisoformat(row["timestamp"])
            event = row["event"]
            if timestamp < interval_start:
                if event == start_event:
                    active = True
                elif event in stop_events:
                    active = False
                continue
            if timestamp >= effective_end:
                break
            if active:
                total += max(0.0, (timestamp - cursor).total_seconds())
            if event == start_event:
                active = True
            elif event in stop_events:
                active = False
            cursor = timestamp
        if active:
            total += max(0.0, (effective_end - cursor).total_seconds())
        return int(round(total))

    def _export_daily_excel_unlocked(self, business_date=None):
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter

        day_text = str(business_date or date.today().isoformat())[:10]
        try:
            report_day = date.fromisoformat(day_text)
        except (TypeError, ValueError) as exc:
            raise ValueError("Report date must use YYYY-MM-DD format.") from exc
        now = datetime.now().astimezone()
        with self._lock:
            events = self.connection.execute(
                """SELECT timestamp, business_date, camera_ip, track_id, event, entered, exited, inside
                   FROM events WHERE branch_id = ? AND business_date = ?
                   ORDER BY timestamp, id""",
                (self.branch_id, day_text),
            ).fetchall()
            health = self.connection.execute(
                """SELECT timestamp, business_date, event, details FROM health_events
                   WHERE branch_id = ? AND business_date <= ?
                   ORDER BY timestamp, id""",
                (self.branch_id, day_text),
            ).fetchall()

        same_day_rows = list(events) + [
            row for row in health if row["business_date"] == day_text
        ]
        report_tz = (
            datetime.fromisoformat(same_day_rows[0]["timestamp"]).tzinfo
            if same_day_rows else now.tzinfo
        )
        events_by_hour = {hour: [] for hour in range(24)}
        for row in events:
            events_by_hour[datetime.fromisoformat(row["timestamp"]).hour].append(row)

        hourly_rows = []
        previous_in = 0
        previous_out = 0
        previous_inside = 0
        for hour in range(24):
            hour_start = datetime.combine(
                report_day, datetime.min.time(), tzinfo=report_tz
            ) + timedelta(hours=hour)
            hour_end = hour_start + timedelta(hours=1)
            is_future = report_day > now.date() or (
                report_day == now.date() and hour_start > now
            )
            hour_events = events_by_hour[hour]
            final_in = int(hour_events[-1]["entered"]) if hour_events else previous_in
            final_out = int(hour_events[-1]["exited"]) if hour_events else previous_out
            closing_inside = (
                int(hour_events[-1]["inside"]) if hour_events else previous_inside
            )
            peak_inside = max(
                [previous_inside] + [int(row["inside"]) for row in hour_events]
            )
            camera_downtime = self._duration_seconds_between(
                health,
                "CAMERA_DISCONNECTED",
                {"CAMERA_CONNECTED", "PROGRAM_STOPPED"},
                hour_start,
                hour_end,
                now,
            )
            program_uptime = self._duration_seconds_between(
                health,
                "PROGRAM_RUNNING",
                {"PROGRAM_STOPPED"},
                hour_start,
                hour_end,
                now,
            )
            hourly_rows.append({
                "hour": f"{hour:02d}:00 - {hour:02d}:59",
                "in": None if is_future else max(0, final_in - previous_in),
                "out": None if is_future else max(0, final_out - previous_out),
                "closing_inside": None if is_future else closing_inside,
                "peak_inside": None if is_future else peak_inside,
                "staff_ignored": None if is_future else sum(
                    "STAFF_" in row["event"] for row in hour_events
                ),
                "unknown_counted": None if is_future else sum(
                    "UNKNOWN_COUNTED" in row["event"] for row in hour_events
                ),
                "camera_downtime": None if is_future else camera_downtime,
                "program_uptime": None if is_future else program_uptime,
            })
            if not is_future:
                previous_in = final_in
                previous_out = final_out
                previous_inside = closing_inside

        wb = Workbook()
        summary = wb.active
        summary.title = "Hourly Report"
        summary.sheet_view.showGridLines = False
        summary.merge_cells("A1:D1")
        summary["A1"] = "DAILY PEOPLE COUNTER REPORT"
        summary["A1"].font = Font(bold=True, size=16, color="FFFFFF")
        summary["A1"].fill = PatternFill("solid", fgColor="1F4E78")
        summary["A1"].alignment = Alignment(horizontal="center")
        summary["A2"] = "Branch"
        summary["B2"] = self.branch_id
        summary["C2"] = "Date"
        summary["D2"] = day_text

        completed_rows = [row for row in hourly_rows if row["in"] is not None]
        total_in = sum(row["in"] for row in completed_rows)
        total_out = sum(row["out"] for row in completed_rows)
        total_staff = sum(row["staff_ignored"] for row in completed_rows)
        totals = (
            ("Total IN", total_in, "Total OUT", total_out),
            ("Staff Ignored", total_staff, "", ""),
        )
        for row_index, values in enumerate(totals, start=4):
            for column, value in enumerate(values, start=1):
                summary.cell(row=row_index, column=column, value=value)
        header_row = 9
        headers = (
            "Hour (24h)", "IN", "OUT", "Staff Ignored",
        )
        for column, value in enumerate(headers, start=1):
            cell = summary.cell(row=header_row, column=column, value=value)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="4472C4")
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for row_index, item in enumerate(hourly_rows, start=header_row + 1):
            values = (
                item["hour"], item["in"], item["out"], item["staff_ignored"],
            )
            for column, value in enumerate(values, start=1):
                summary.cell(row=row_index, column=column, value=value)
        summary.freeze_panes = f"A{header_row + 1}"
        summary.auto_filter.ref = f"A{header_row}:D{header_row + len(hourly_rows)}"
        widths = (20, 14, 14, 18)
        for index, width in enumerate(widths, start=1):
            summary.column_dimensions[get_column_letter(index)].width = width

        self.reports_dir.mkdir(parents=True, exist_ok=True)
        target = self.reports_dir / f"{day_text}.xlsx"
        temporary = target.with_suffix(".tmp.xlsx")
        wb.save(temporary)
        wb.close()
        os.replace(temporary, target)
        return target

    def run_maintenance(self):
        cutoff = (date.today() - timedelta(days=self.event_retention_days)).isoformat()
        with self._lock:
            self.connection.execute(
                "DELETE FROM events WHERE branch_id = ? AND business_date < ?",
                (self.branch_id, cutoff),
            )
            self.connection.execute(
                "DELETE FROM health_events WHERE branch_id = ? AND business_date < ?",
                (self.branch_id, cutoff),
            )
            self.connection.commit()
            self.connection.execute("PRAGMA wal_checkpoint(TRUNCATE)")

        report_cutoff = date.today() - timedelta(days=self.report_retention_days)
        for report in self.reports_dir.glob("????-??-??.xlsx"):
            try:
                report_day = date.fromisoformat(report.stem)
                if report_day < report_cutoff:
                    report.unlink()
            except (OSError, ValueError):
                continue

    def close(self, state=None, export_today=True):
        if self._closed:
            return
        self._report_stop.set()
        self._report_requested.set()
        self._report_thread.join(timeout=self.report_refresh_delay_s + 10.0)
        if state is not None:
            self.save_state(state)
        if export_today:
            self.export_daily_excel(
                state["business_date"] if state else date.today().isoformat()
            )
        with self._lock:
            self.connection.close()
            self._closed = True


class HeartbeatWriter:
    def __init__(self, filename):
        self.filename = Path(filename)
        self.filename.parent.mkdir(parents=True, exist_ok=True)
        self._lock = threading.Lock()

    def write(self, **payload):
        data = {
            "timestamp": datetime.now().astimezone().isoformat(timespec="seconds"),
            **payload,
        }
        temporary = self.filename.with_suffix(self.filename.suffix + ".tmp")
        with self._lock:
            temporary.write_text(
                json.dumps(data, ensure_ascii=False, indent=2) + "\n",
                encoding="utf-8",
            )
            os.replace(temporary, self.filename)
        return data
