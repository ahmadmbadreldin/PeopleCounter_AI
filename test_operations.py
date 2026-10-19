import io
import json
import logging
import os
import tempfile
import time
import unittest
from contextlib import redirect_stdout
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

import status
from runtime_logging import configure_rotating_logger
from runtime_storage import SQLiteStore, empty_state
from watchdog import _stop_child, heartbeat_is_stale


class ExcelAndDatabaseTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        root = Path(self.temp.name)
        self.store = SQLiteStore(
            root / "data" / "branch.sqlite3",
            "Branch_A",
            "10.0.0.7",
            root / "reports",
            report_refresh_delay_s=0.25,
        )
        self.root = root

    def tearDown(self):
        self.store.close(export_today=False)
        self.temp.cleanup()

    def test_sqlite_uses_wal_and_expected_tables(self):
        journal = self.store.connection.execute("PRAGMA journal_mode").fetchone()[0]
        tables = {
            row[0] for row in self.store.connection.execute(
                "SELECT name FROM sqlite_master WHERE type='table'"
            )
        }
        self.assertEqual(journal.lower(), "wal")
        self.assertTrue({"counter_state", "events", "health_events"}.issubset(tables))

    def test_event_and_state_are_committed_together(self):
        state = empty_state("2026-10-17")
        state.update(total_in=2, total_out=1, peak_inside=1)
        self.store.add_event(22, "IN", state)
        event = self.store.connection.execute("SELECT * FROM events").fetchone()
        saved = self.store.connection.execute("SELECT * FROM counter_state").fetchone()
        self.assertEqual(event["inside"], 1)
        self.assertEqual(saved["total_in"], 2)

    def test_daily_excel_contains_one_row_per_hour(self):
        state = empty_state("2026-10-17")
        state.update(total_in=3, total_out=1, peak_inside=2)
        with patch.object(
            self.store, "_now", return_value="2026-10-17T09:15:00+02:00"
        ):
            self.store.add_event(1, "IN_UNKNOWN_COUNTED", state)
        with self.store._lock:
            self.store.connection.executemany(
                """INSERT INTO health_events
                   (timestamp, business_date, branch_id, event, details)
                   VALUES (?, ?, ?, ?, ?)""",
                (
                    ("2026-10-17T09:00:00+02:00", "2026-10-17", "Branch_A", "PROGRAM_RUNNING", ""),
                    ("2026-10-17T17:00:00+02:00", "2026-10-17", "Branch_A", "PROGRAM_STOPPED", ""),
                    ("2026-10-17T10:00:00+02:00", "2026-10-17", "Branch_A", "CAMERA_DISCONNECTED", "cable"),
                    ("2026-10-17T10:05:20+02:00", "2026-10-17", "Branch_A", "CAMERA_CONNECTED", ""),
                ),
            )
            self.store.connection.commit()
        report = self.store.export_daily_excel("2026-10-17")
        workbook = load_workbook(report, read_only=True, data_only=True)
        try:
            self.assertEqual(report.name, "2026-10-17.xlsx")
            self.assertEqual(workbook.sheetnames, ["Hourly Report"])
            sheet = workbook["Hourly Report"]
            rows = {row[0]: row for row in sheet.iter_rows(min_row=10, values_only=True)}
            self.assertEqual(len(rows), 24)
            self.assertEqual(
                tuple(cell.value for cell in sheet[9]),
                ("Hour (24h)", "IN", "OUT", "Staff Ignored"),
            )
            self.assertEqual(rows["00:00 - 00:59"][1:4], (0, 0, 0))
            hour_09 = rows["09:00 - 09:59"]
            self.assertEqual(hour_09[1:4], (3, 1, 0))
            self.assertEqual(sheet["B4"].value, 3)
            self.assertEqual(sheet["D4"].value, 1)
            self.assertEqual(sheet["B5"].value, 0)
        finally:
            workbook.close()

    def test_health_details_are_bounded(self):
        self.store.log_health("ERROR", "x" * 5000, "2026-10-17")
        details = self.store.connection.execute(
            "SELECT details FROM health_events ORDER BY id DESC"
        ).fetchone()[0]
        self.assertEqual(len(details), 2000)

    def test_retention_removes_old_database_rows_and_reports(self):
        old_day = "2000-01-01"
        state = empty_state(old_day)
        self.store.add_event(1, "IN", state)
        self.store.log_health("OLD", "old", old_day)
        old_report = self.store.export_daily_excel(old_day)
        self.assertTrue(old_report.exists())
        self.store.run_maintenance()
        self.assertEqual(self.store.connection.execute("SELECT COUNT(*) FROM events").fetchone()[0], 0)
        self.assertEqual(self.store.connection.execute("SELECT COUNT(*) FROM health_events").fetchone()[0], 0)
        self.assertFalse(old_report.exists())

    def test_close_is_idempotent(self):
        self.store.close(export_today=False)
        self.store.close(export_today=False)


class StatusAndWatchdogTests(unittest.TestCase):
    def test_status_command_reports_healthy_camera(self):
        with tempfile.TemporaryDirectory() as folder:
            heartbeat = Path(folder) / "heartbeat.json"
            heartbeat.write_text(json.dumps({
                "status": "running", "timestamp": "now",
                "camera_connected": True, "entered": 5,
                "exited": 2, "inside": 3,
            }), encoding="utf-8")
            config = {"branch_id": "B1", "runtime": {"watchdog_stale_s": 45}}
            output = io.StringIO()
            with patch.object(status, "load_branch_config", return_value=config), \
                 patch.object(status, "resolve_project_path", return_value=heartbeat), \
                 redirect_stdout(output):
                code = status.main()
            self.assertEqual(code, 0)
            self.assertIn("Program: RUNNING", output.getvalue())
            self.assertIn("Camera: CONNECTED", output.getvalue())

    def test_status_handles_missing_and_corrupt_heartbeat(self):
        with tempfile.TemporaryDirectory() as folder:
            heartbeat = Path(folder) / "heartbeat.json"
            config = {"branch_id": "B1", "runtime": {}}
            with patch.object(status, "load_branch_config", return_value=config), \
                 patch.object(status, "resolve_project_path", return_value=heartbeat), \
                 redirect_stdout(io.StringIO()):
                self.assertEqual(status.main(), 1)
                heartbeat.write_text("bad json", encoding="utf-8")
                self.assertEqual(status.main(), 1)

    def test_stale_heartbeat_detection(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "hb.json"
            self.assertTrue(heartbeat_is_stale(path, 10))
            path.write_text("{}", encoding="utf-8")
            self.assertFalse(heartbeat_is_stale(path, 10))
            old = time.time() - 60
            os.utime(path, (old, old))
            self.assertTrue(heartbeat_is_stale(path, 10))

    def test_watchdog_terminates_then_kills_hung_child(self):
        class Process:
            pid = 7
            killed = False
            terminated = False

            def poll(self):
                return None

            def terminate(self):
                self.terminated = True

            def wait(self, timeout):
                import subprocess
                if not self.killed:
                    raise subprocess.TimeoutExpired("fake", timeout)
                return 1

            def kill(self):
                self.killed = True

        process = Process()
        logger = logging.getLogger("watchdog-test")
        _stop_child(process, logger)
        self.assertTrue(process.terminated)
        self.assertTrue(process.killed)


class LogRotationTests(unittest.TestCase):
    def test_log_rotates_and_is_bounded(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "runtime.log"
            name = f"rotation-test-{time.time_ns()}"
            logger = configure_rotating_logger(name, path, max_bytes=120, backup_count=2)
            for index in range(50):
                logger.info("line %s %s", index, "x" * 50)
            for handler in logger.handlers:
                handler.flush()
                handler.close()
            self.assertTrue(path.exists())
            self.assertLessEqual(len(list(path.parent.glob("runtime.log*"))), 3)


class InstallationArtifactTests(unittest.TestCase):
    def test_branch_installer_contains_required_safety_and_startup_actions(self):
        root = Path(__file__).resolve().parent
        script = (root / "branch_windows_setup.ps1").read_text(encoding="utf-8")
        self.assertIn("New-ScheduledTaskTrigger -AtLogOn", script)
        self.assertIn("standby-timeout-ac 0", script)
        self.assertIn("hibernate-timeout-ac 0", script)
        self.assertIn("standby-timeout-dc 0", script)
        self.assertIn("MultipleInstances IgnoreNew", script)
        self.assertIn("watchdog.py", script)

    def test_all_operator_launchers_exist(self):
        root = Path(__file__).resolve().parent
        for name in (
            "1 - Configure Branch.cmd",
            "2 - Start Monitoring.cmd",
            "Start People Counter.cmd",
            "Install Branch PC.cmd",
            "Run People Counter 24-7.cmd",
            "Check People Counter Status.cmd",
            "Remove Branch Startup.cmd",
            "Stop Old People Counter.cmd",
        ):
            with self.subTest(name=name):
                self.assertTrue((root / name).is_file())

    def test_configuration_and_monitoring_launchers_are_separate(self):
        root = Path(__file__).resolve().parent
        configure = (root / "1 - Configure Branch.cmd").read_text(encoding="utf-8")
        monitor = (root / "2 - Start Monitoring.cmd").read_text(encoding="utf-8")
        self.assertIn("setup_branch.py --configure-only", configure)
        self.assertNotIn("watchdog.py", configure)
        self.assertIn("watchdog.py", monitor)
        self.assertNotIn("setup_branch.py", monitor)
        self.assertIn("validate_config", monitor)

    def test_stop_script_is_scoped_to_this_project(self):
        root = Path(__file__).resolve().parent
        script = (root / "stop_people_counter.ps1").read_text(encoding="utf-8")
        self.assertIn("[Regex]::Escape($projectRoot", script)
        self.assertIn("people_counter_hikvision.py", script)
        self.assertIn("watchdog.py", script)
        self.assertIn("Stop-ScheduledTask", script)
        self.assertNotIn("Stop-Process -Name python", script)


if __name__ == "__main__":
    unittest.main()
